"""Build a price-update CSV from a PSA "Price Situationer of Selected Agricultural Commodities" workbook.

Usage (from backend/; needs `pip install openpyxl` on the machine doing the update, not for the app):
    python data/price_updates/build_psa_update.py <statistical_tables.xlsx> <out.csv> \
        --phase "Second Phase Sep 2026" --observed-at 2026-09-17 --release-url <PSA page URL>
Then import it:
    python ds_ds_update_prices.py <out.csv> --allow-big-changes      (after reviewing the flagged changes)

Reads the NCR row of each table for the requested phase. Only commodities whose PSA specification
matches one of our ingredients are mapped (see MAPPING); everything else is ignored, never guessed.
"""
import argparse
import csv
import re
import warnings

warnings.filterwarnings("ignore")

# PSA commodity (exact table header) -> (ingredient_id, unit the PSA price is quoted in)
MAPPING = {
    "RICE, REGULAR MILLED": ("rice", "kg"),
    "PORK, KASIM": ("pork_kasim", "kg"),
    "DRESSED CHICKEN": ("chicken", "kg"),
    "CHICKEN EGG, MED. (per piece)": ("eggs", "piece"),
    "TILAPIA": ("tilapia", "kg"),
    "AMPALAYA": ("ampalaya", "kg"),
    "CABBAGE": ("cabbage", "kg"),
    "CARROT": ("carrot", "kg"),
    "EGGPLANT": ("talong", "kg"),
    "PECHAY, NATIVE": ("pechay", "kg"),
    "STRING BEANS": ("sitaw", "kg"),
    "TOMATO": ("tomato", "kg"),
    "WHITE POTATO": ("potato", "kg"),
    "ONION, RED": ("onion", "kg"),
    "GARLIC, IMPORTED": ("garlic", "kg"),
    "GINGER, HAWAIIAN": ("ginger", "kg"),
    "CALAMANSI": ("calamansi", "kg"),
    "BANANA, LAKATAN": ("lakatan", "kg"),
    "MANGO, CARABAO": ("mango", "kg"),
    "BROWN SUGAR (in PhP per Kilogram)": ("sugar", "kg"),
    # Deliberately NOT mapped: "COOKING OIL (in PhP per Liter)" has no type/size specification, while our
    # oil row is the DA-specified palm oil 350 ml that shoppers buy in small bottles. Not the same product.
}
# minimum amount a vendor sells, in the PSA unit (same values as the seed data)
MIN_QTY = {"kg": "0.25", "piece": "1"}
MIN_QTY_BY_ID = {"garlic": "0.05", "ginger": "0.05", "onion": "0.1", "tomato": "0.1", "calamansi": "0.1"}

norm = lambda s: re.sub(r"\s+", " ", str(s or "")).strip()


def extract(xlsx_path, phase):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    found = {}
    for ws in wb.worksheets:
        if not ws.title.startswith("Table ") or ws.title.startswith("Table_All"):
            continue
        rows = list(ws.iter_rows(values_only=True))
        h = next(i for i, r in enumerate(rows) if norm(r[0]).startswith("Region / Province"))
        ncr = next(r for r in rows if norm(r[0]).startswith("NCR"))
        commodity = None
        for j, v in enumerate(rows[h]):
            if j and norm(v):
                commodity = norm(v)
            if j and commodity and norm(rows[h + 1][j]) == phase:
                found[commodity] = (ncr[j], ws.title, j)
    return found


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx"); ap.add_argument("out_csv")
    ap.add_argument("--phase", required=True); ap.add_argument("--observed-at", required=True)
    ap.add_argument("--release-url", required=True)
    a = ap.parse_args()
    import openpyxl.utils as u
    found = extract(a.xlsx, a.phase)
    with open(a.out_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["ingredient", "amount_php", "quantity", "unit", "observed_at", "source_type",
                    "source_reference", "market_or_area", "minimum_purchase_quantity"])
        for commodity, (iid, unit) in MAPPING.items():
            if commodity not in found:
                print(f"  skip {commodity}: not in workbook"); continue
            value, sheet, j = found[commodity]
            if not isinstance(value, (int, float)):
                print(f"  skip {commodity}: NCR value is {value!r}"); continue
            ref = (f"PSA Price Situationer of Selected Agricultural Commodities, {a.phase} "
                   f"(statistical tables, sheet '{sheet}', column {u.get_column_letter(j + 1)}, NCR row, "
                   f"'{commodity}'). {a.release_url}")
            w.writerow([iid, round(float(value), 2), 1, unit, a.observed_at, "official_reference", ref,
                        "NCR average (PSA)", MIN_QTY_BY_ID.get(iid, MIN_QTY[unit])])
            print(f"  {iid:<12} PHP {float(value):8.2f} per 1 {unit:<5} <- {commodity}")


if __name__ == "__main__":
    main()
